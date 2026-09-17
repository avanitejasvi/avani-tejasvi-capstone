import csv
import io
import json
import re
import sys
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Optional

from dotenv import load_dotenv
import os

from pydantic import BaseModel, Field

from openpyxl import load_workbook

from google.oauth2.credentials import Credentials
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError
from googleapiclient.http import MediaIoBaseDownload

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
DRIVE_FOLDER_NAME = "meal-menu-match"
DRIVE_FILE_NAME = "weekly_menu_google_sheets_ready.xlsx"
GOOGLE_SHEET_MIME_TYPE = "application/vnd.google-apps.spreadsheet"
EAT_THRESHOLD = 3.0
GOOGLE_SCOPES = [
    "https://www.googleapis.com/auth/drive.readonly",
    "https://www.googleapis.com/auth/calendar",
]

# The real Drive sheet is a weekly grid: one column per weekday, one row per
# dish category, grouped into meal-slot sections by header rows like these.
GRID_SECTION_HEADERS = {"breakfast", "lunch", "evening snacks", "dinner"}
GRID_EMPTY_MARKERS = {"n/a", "n/l", "na", "nil", "-", ""}


class GoogleAuthUnavailable(Exception):
    pass


class KnownDish(BaseModel):
    name: str
    tags: list[str]
    rating: int
    last_seen: date


class UserPreferences(BaseModel):
    user_id: str
    dietary_restrictions: list[str]
    skip_meal_slots: list[str]
    known_dishes: list[KnownDish]
    tag_weights: dict[str, float] = Field(default_factory=dict)


class MenuArchiveItem(BaseModel):
    id: str
    name: str
    tags: list[str]
    first_seen_date: date


class ScoredDish(BaseModel):
    name: str
    score: float
    source: str
    decision: str


class MealEvent(BaseModel):
    meal_slot: str
    event_date: date
    scored_items: list[ScoredDish]
    calendar_event_id: Optional[str] = None
    response: Optional[str] = None


@dataclass
class Trace:
    step: int
    thought: str
    action: str
    action_input: dict
    observation: str


def _load_google_credentials() -> Credentials:
    client_id = os.getenv("GOOGLE_CLIENT_ID")
    client_secret = os.getenv("GOOGLE_CLIENT_SECRET")
    refresh_token = os.getenv("GOOGLE_REFRESH_TOKEN")
    if not (client_id and client_secret and refresh_token):
        raise GoogleAuthUnavailable(
            "Missing GOOGLE_CLIENT_ID / GOOGLE_CLIENT_SECRET / GOOGLE_REFRESH_TOKEN "
            "in the environment; run the OAuth consent flow once and store the "
            "resulting refresh token in .env before live Drive/Calendar calls will work."
        )
    return Credentials(
        token=None,
        refresh_token=refresh_token,
        token_uri="https://oauth2.googleapis.com/token",
        client_id=client_id,
        client_secret=client_secret,
        scopes=GOOGLE_SCOPES,
    )


def build_drive_service():
    return build("drive", "v3", credentials=_load_google_credentials(), cache_discovery=False)


def build_calendar_service():
    return build("calendar", "v3", credentials=_load_google_credentials(), cache_discovery=False)


class FilesystemTool:
    @staticmethod
    def read_json(path: Path) -> dict:
        return json.loads(path.read_text())

    @staticmethod
    def write_json(path: Path, payload: dict) -> None:
        path.write_text(json.dumps(payload, indent=2, default=str))

    @staticmethod
    def read_csv_rows(path: Path) -> list[dict]:
        with path.open(newline="") as handle:
            return list(csv.DictReader(handle))


class GoogleDriveSkill:
    def __init__(self, folder_name: str = DRIVE_FOLDER_NAME, file_name: str = DRIVE_FILE_NAME):
        self.folder_name = folder_name
        self.file_name = file_name

    def _find_file(self, service) -> tuple[str, str]:
        folder_results = (
            service.files()
            .list(
                q=f"name='{self.folder_name}' and mimeType='application/vnd.google-apps.folder' and trashed=false",
                fields="files(id,name)",
            )
            .execute()
        )
        folders = folder_results.get("files", [])
        if not folders:
            raise FileNotFoundError(f"No Drive folder named '{self.folder_name}' found")
        folder_id = folders[0]["id"]

        file_results = (
            service.files()
            .list(
                q=f"name='{self.file_name}' and '{folder_id}' in parents and trashed=false",
                fields="files(id,name,mimeType)",
            )
            .execute()
        )
        files = file_results.get("files", [])
        if not files:
            raise FileNotFoundError(f"No file named '{self.file_name}' found in Drive folder '{self.folder_name}'")
        return files[0]["id"], files[0]["mimeType"]

    def read_menu_sheet(self, meal_slot: str, event_date: date) -> list[MenuArchiveItem]:
        service = build_drive_service()
        file_id, mime_type = self._find_file(service)
        is_native_sheet = mime_type == GOOGLE_SHEET_MIME_TYPE

        request = (
            service.files().export_media(fileId=file_id, mimeType="text/csv")
            if is_native_sheet
            else service.files().get_media(fileId=file_id)
        )
        buffer = io.BytesIO()
        downloader = MediaIoBaseDownload(buffer, request)
        done = False
        while not done:
            _, done = downloader.next_chunk()
        buffer.seek(0)

        rows = (
            list(csv.DictReader(io.StringIO(buffer.read().decode("utf-8"))))
            if is_native_sheet
            else _read_xlsx_rows(buffer)
        )
        return _parse_weekly_grid(rows, meal_slot, event_date)


def _read_xlsx_rows(buffer: io.BytesIO) -> list[dict]:
    workbook = load_workbook(buffer, data_only=True)
    sheet = workbook.active
    rows_iter = sheet.iter_rows(values_only=True)
    header = [str(cell).strip() if cell is not None else "" for cell in next(rows_iter)]
    records = []
    for row in rows_iter:
        if all(cell is None for cell in row):
            continue
        records.append({header[i]: row[i] for i in range(len(header)) if i < len(row)})
    return records


def _find_day_column(rows: list[dict], event_date: date) -> Optional[str]:
    if not rows:
        return None
    date_row = rows[0]
    for column, raw_value in date_row.items():
        if not column or column.strip().lower() == "day" or not raw_value:
            continue
        try:
            parsed = datetime.strptime(f"{str(raw_value).strip()}-{event_date.year}", "%d-%b-%Y").date()
        except ValueError:
            continue
        if parsed == event_date:
            return column
    return None


def _parse_weekly_grid(rows: list[dict], meal_slot: str, event_date: date) -> list[MenuArchiveItem]:
    """Parses the real 'meal-menu-match' sheet: one column per weekday (row 0 maps
    each column to an actual date), one row per dish category, grouped into
    meal-slot sections by header rows (Breakfast/Lunch/Evening Snacks/Dinner)."""
    day_column = _find_day_column(rows, event_date)
    if day_column is None:
        return []

    items: list[MenuArchiveItem] = []
    current_section: Optional[str] = None
    for index, row in enumerate(rows[1:], start=1):
        label = str(row.get("Day") or "").strip()
        if not label:
            continue
        if label.lower() in GRID_SECTION_HEADERS:
            current_section = label.lower()
            continue
        if current_section != meal_slot:
            continue

        dish_name = row.get(day_column)
        if not dish_name or str(dish_name).strip().lower() in GRID_EMPTY_MARKERS:
            continue
        dish_name = str(dish_name).strip()
        tags = [tag for tag in re.split(r"[\s/\-]+", label.lower()) if tag]

        items.append(
            MenuArchiveItem(
                id=f"{event_date.isoformat()}-{meal_slot}-{index}",
                name=dish_name,
                tags=tags,
                first_seen_date=event_date,
            )
        )
    return items


class GoogleCalendarSkill:
    def __init__(self, calendar_id: str = "primary"):
        self.calendar_id = calendar_id

    def create_meal_event(self, meal_slot: str, event_date: date, description: str) -> str:
        service = build_calendar_service()
        start_hour = {"breakfast": 8, "lunch": 13, "dinner": 20}.get(meal_slot, 13)
        start_dt = datetime.combine(event_date, datetime.min.time()).replace(hour=start_hour)
        end_dt = start_dt.replace(hour=start_hour + 1)
        body = {
            "summary": f"Mess meal check: {meal_slot.title()}",
            "description": description,
            "start": {"dateTime": start_dt.isoformat(), "timeZone": "Asia/Kolkata"},
            "end": {"dateTime": end_dt.isoformat(), "timeZone": "Asia/Kolkata"},
        }
        created = service.events().insert(calendarId=self.calendar_id, body=body).execute()
        return created["id"]

    def get_response(self, event_id: str) -> Optional[str]:
        service = build_calendar_service()
        event = service.events().get(calendarId=self.calendar_id, eventId=event_id).execute()
        return "no" if event.get("status") == "cancelled" else None

    def delete_event(self, event_id: str) -> None:
        build_calendar_service().events().delete(calendarId=self.calendar_id, eventId=event_id).execute()


class MenuPreferenceMatchingSkill:
    def compute_tag_weights(self, prefs: UserPreferences) -> dict[str, float]:
        tag_ratings: dict[str, list[int]] = {}
        for dish in prefs.known_dishes:
            for tag in dish.tags:
                tag_ratings.setdefault(tag, []).append(dish.rating)
        return {tag: sum(ratings) / len(ratings) for tag, ratings in tag_ratings.items()}

    def is_slot_skipped(self, prefs: UserPreferences, meal_slot: str) -> bool:
        return meal_slot in prefs.skip_meal_slots

    def violates_dietary_restriction(self, item: MenuArchiveItem, prefs: UserPreferences) -> bool:
        item_tags = {tag.lower() for tag in item.tags}
        item_name = item.name.lower()
        return any(
            restriction.lower() in item_tags or restriction.lower() in item_name
            for restriction in prefs.dietary_restrictions
        )

    def score_item(self, item: MenuArchiveItem, prefs: UserPreferences) -> ScoredDish:
        for dish in prefs.known_dishes:
            if dish.name.strip().lower() == item.name.strip().lower():
                decision = "eat" if dish.rating >= EAT_THRESHOLD else "skip"
                return ScoredDish(name=item.name, score=float(dish.rating), source="known_dish", decision=decision)
        recognized = [prefs.tag_weights[tag] for tag in item.tags if tag in prefs.tag_weights]
        if not recognized:
            return ScoredDish(name=item.name, score=0.0, source="no_data", decision="skip")
        estimate = sum(recognized) / len(recognized)
        decision = "eat" if estimate >= EAT_THRESHOLD else "skip"
        return ScoredDish(name=item.name, score=round(estimate, 2), source="tag_weight_estimate", decision=decision)

    def register_new_dishes(
        self,
        prefs: UserPreferences,
        scored: list[ScoredDish],
        item_tags_by_name: dict[str, list[str]],
        event_date: date,
    ) -> UserPreferences:
        existing = {dish.name.strip().lower() for dish in prefs.known_dishes}
        for result in scored:
            key = result.name.strip().lower()
            if key in existing or result.source != "tag_weight_estimate":
                continue
            rating = max(1, min(5, round(result.score)))
            prefs.known_dishes.append(
                KnownDish(name=result.name, tags=item_tags_by_name.get(result.name, []), rating=rating, last_seen=event_date)
            )
        return prefs

    def apply_decline_feedback(self, prefs: UserPreferences, scored: list[ScoredDish]) -> UserPreferences:
        scored_names = {s.name.strip().lower() for s in scored}
        for dish in prefs.known_dishes:
            if dish.name.strip().lower() in scored_names:
                dish.rating = max(1, dish.rating - 1)
        return prefs

    def build_event_description(self, scored: list[ScoredDish]) -> str:
        known_favorites = [s for s in scored if s.source == "known_dish" and s.decision == "eat"]
        new_estimated = [s for s in scored if s.source == "tag_weight_estimate" and s.decision == "eat"]
        no_data = [s for s in scored if s.source == "no_data"]
        skipped = [s for s in scored if s.decision == "skip" and s.source != "no_data"]
        lines = []
        if known_favorites:
            lines.append("Known favorites: " + ", ".join(f"{s.name} ({s.score:.1f})" for s in known_favorites))
        if new_estimated:
            lines.append("New, estimated matches: " + ", ".join(f"{s.name} (~{s.score:.1f})" for s in new_estimated))
        if no_data:
            lines.append("New dish, no data yet: " + ", ".join(s.name for s in no_data))
        if skipped:
            lines.append("Not worth it today: " + ", ".join(s.name for s in skipped))
        if not lines:
            lines.append("No menu items available to evaluate.")
        return "\n".join(lines)


class ReActMealAgent:
    def __init__(
        self,
        user_id: str,
        drive_skill: GoogleDriveSkill,
        calendar_skill: GoogleCalendarSkill,
        matcher: MenuPreferenceMatchingSkill,
        fs: FilesystemTool,
    ):
        self.user_id = user_id
        self.drive = drive_skill
        self.calendar = calendar_skill
        self.matcher = matcher
        self.fs = fs
        self.trace: list[Trace] = []
        self._step_counter = 0

    def _log(self, thought: str, action: str, action_input: dict, observation: str) -> None:
        self._step_counter += 1
        entry = Trace(self._step_counter, thought, action, action_input, observation)
        self.trace.append(entry)
        print(f"\n--- Step {entry.step} ---")
        print(f"Thought: {entry.thought}")
        print(f"Action: {entry.action}")
        print(f"Action Input: {entry.action_input}")
        print(f"Observation: {entry.observation}")

    def _preferences_path(self) -> Path:
        return DATA_DIR / f"preferences_{self.user_id}.json"

    def perceive_preferences(self) -> UserPreferences:
        path = self._preferences_path()
        raw = self.fs.read_json(path)
        prefs = UserPreferences(**raw)
        self._log(
            "I need the student's stored preferences before judging any dish.",
            "filesystem.read_json",
            {"path": str(path)},
            f"Loaded {len(prefs.known_dishes)} known dish(es), "
            f"{len(prefs.dietary_restrictions)} dietary restriction(s), skip_meal_slots={prefs.skip_meal_slots}.",
        )
        prefs.tag_weights = self.matcher.compute_tag_weights(prefs)
        self._log(
            "tag_weights is derived data recomputed from known_dishes every run, never hand-maintained.",
            "matcher.compute_tag_weights",
            {"known_dishes": len(prefs.known_dishes)},
            f"Recomputed {len(prefs.tag_weights)} tag weight(s): {prefs.tag_weights}",
        )
        return prefs

    def reason_skip_check(self, prefs: UserPreferences, meal_slot: str) -> bool:
        skip = self.matcher.is_slot_skipped(prefs, meal_slot)
        self._log(
            f"Before scoring anything, check whether '{meal_slot}' is in skip_meal_slots — "
            "that's a behavior rule, not a taste signal.",
            "matcher.is_slot_skipped",
            {"meal_slot": meal_slot},
            f"'{meal_slot}' {'IS' if skip else 'is NOT'} in skip_meal_slots={prefs.skip_meal_slots}.",
        )
        return skip

    def act_perceive_menu(self, meal_slot: str, event_date: date) -> list[MenuArchiveItem]:
        drive_action = "google_drive.read_menu_sheet"
        drive_input = {"meal_slot": meal_slot, "folder": self.drive.folder_name, "file": self.drive.file_name}
        try:
            items = self.drive.read_menu_sheet(meal_slot, event_date)
            self._log(
                "Try the real Drive source first — SKILL.md requires the live "
                "'meal-menu-match' sheet, not a hand-typed menu.",
                drive_action,
                drive_input,
                f"Drive returned {len(items)} item(s) for slot '{meal_slot}'.",
            )
            return items
        except (GoogleAuthUnavailable, HttpError, FileNotFoundError) as exc:
            self._log(
                "Try the real Drive source first — SKILL.md requires the live "
                "'meal-menu-match' sheet, not a hand-typed menu.",
                drive_action,
                drive_input,
                f"Drive read failed ({exc}); falling back to the local sample menu, "
                "matching SKILL.md's test-before-connecting build order.",
            )
        path = DATA_DIR / "menu_archive_sample.csv"
        rows = self.fs.read_csv_rows(path)
        items = [
            MenuArchiveItem(id=row["id"], name=row["name"], tags=row["tags"].split("|"), first_seen_date=row["first_seen_date"])
            for row in rows
            if row["meal_slot"] == meal_slot
        ]
        self._log(
            "Local fixture stands in for the Drive sheet until credentials are configured.",
            "filesystem.read_csv_rows",
            {"path": str(path)},
            f"Loaded {len(items)} local item(s) for slot '{meal_slot}'.",
        )
        return items

    def act_score_items(self, items: list[MenuArchiveItem], prefs: UserPreferences) -> list[ScoredDish]:
        scored = []
        for item in items:
            if self.matcher.violates_dietary_restriction(item, prefs):
                self._log(
                    f"Hard-exclude check for '{item.name}' runs before any scoring.",
                    "matcher.violates_dietary_restriction",
                    {"item": item.name, "tags": item.tags},
                    f"'{item.name}' excluded — violates dietary_restrictions={prefs.dietary_restrictions}.",
                )
                continue
            result = self.matcher.score_item(item, prefs)
            self._log(
                f"'{item.name}' — check known_dishes by exact name first; only fall back "
                "to tag_weights if it's genuinely new.",
                "matcher.score_item",
                {"item": item.name, "tags": item.tags},
                f"source={result.source} score={result.score} decision={result.decision}",
            )
            scored.append(result)
        return scored

    def act_register_new_dishes(
        self, prefs: UserPreferences, scored: list[ScoredDish], items: list[MenuArchiveItem], event_date: date
    ) -> UserPreferences:
        item_tags_by_name = {item.name: item.tags for item in items}
        before = len(prefs.known_dishes)
        prefs = self.matcher.register_new_dishes(prefs, scored, item_tags_by_name, event_date)
        self._log(
            "A tag-weight estimate joins known_dishes now so tag_weights self-improves next run; "
            "a no-data dish stays unrated rather than getting a silently guessed score.",
            "matcher.register_new_dishes",
            {"scored_count": len(scored)},
            f"Added {len(prefs.known_dishes) - before} new dish(es) to known_dishes.",
        )
        return prefs

    def act_build_description(self, scored: list[ScoredDish]) -> str:
        description = self.matcher.build_event_description(scored)
        self._log(
            "Compose the invite body, separating confident matches from estimates so it stays honest.",
            "matcher.build_event_description",
            {"scored_count": len(scored)},
            description.replace("\n", " | "),
        )
        return description

    def act_create_event(self, meal_slot: str, event_date: date, description: str) -> Optional[str]:
        action_input = {"meal_slot": meal_slot, "event_date": str(event_date)}
        try:
            event_id = self.calendar.create_meal_event(meal_slot, event_date, description)
            self._log(
                "Push the result into Calendar as the actual interface, per SKILL.md — no separate custom screen.",
                "google_calendar.create_meal_event",
                action_input,
                f"Created calendar event {event_id}.",
            )
            return event_id
        except (GoogleAuthUnavailable, HttpError) as exc:
            self._log(
                "Push the result into Calendar as the actual interface, per SKILL.md — no separate custom screen.",
                "google_calendar.create_meal_event",
                action_input,
                f"Calendar unavailable ({exc}); recommendation is computed but not pushed to a live calendar this run.",
            )
            return None

    def act_check_response(self, event_id: Optional[str]) -> Optional[str]:
        if event_id is None:
            self._log(
                "No live event exists to poll.",
                "google_calendar.get_response",
                {"event_id": None},
                "Skipped — nothing to check.",
            )
            return None
        try:
            response = self.calendar.get_response(event_id)
        except (GoogleAuthUnavailable, HttpError) as exc:
            self._log(
                "Ask Calendar whether the student accepted or declined the invite.",
                "google_calendar.get_response",
                {"event_id": event_id},
                f"Could not read RSVP ({exc}); treating as pending.",
            )
            return None
        self._log(
            "Ask Calendar whether the student accepted or declined the invite.",
            "google_calendar.get_response",
            {"event_id": event_id},
            f"RSVP status: {response or 'pending'}.",
        )
        return response

    def act_handle_decline(
        self, prefs: UserPreferences, scored: list[ScoredDish], event_id: Optional[str]
    ) -> UserPreferences:
        prefs = self.matcher.apply_decline_feedback(prefs, scored)
        self._log(
            "Student declined — nudge the estimated ratings down before discarding the event, "
            "so the signal isn't lost.",
            "matcher.apply_decline_feedback",
            {"event_id": event_id},
            "Lowered ratings for the evaluated dishes by 1 (floor 1).",
        )
        if event_id is not None:
            try:
                self.calendar.delete_event(event_id)
                observation = "Event deleted."
            except (GoogleAuthUnavailable, HttpError) as exc:
                observation = f"Delete failed ({exc}) — no live calendar this run."
            self._log(
                "Delete the event now that its feedback has been captured.",
                "google_calendar.delete_event",
                {"event_id": event_id},
                observation,
            )
        return prefs

    def act_persist_preferences(self, prefs: UserPreferences) -> None:
        path = self._preferences_path()
        payload = json.loads(prefs.model_dump_json())
        payload["tag_weights"] = {}
        self.fs.write_json(path, payload)
        self._log(
            "known_dishes may have changed this run (new dishes and/or decline feedback); "
            "tag_weights is never persisted since it is derived fresh every run.",
            "filesystem.write_json",
            {"path": str(path)},
            f"Wrote {len(prefs.known_dishes)} known dish(es) back to disk.",
        )

    def run(self, meal_slot: str, event_date: date, simulated_response: Optional[str] = None) -> MealEvent:
        prefs = self.perceive_preferences()

        if self.reason_skip_check(prefs, meal_slot):
            self._log(
                "Skip rule matched — halt before any scoring or Drive/Calendar calls.",
                "workflow.halt",
                {"meal_slot": meal_slot},
                "No event created; this meal slot is excluded by user preference.",
            )
            return MealEvent(meal_slot=meal_slot, event_date=event_date, scored_items=[], calendar_event_id=None, response=None)

        items = self.act_perceive_menu(meal_slot, event_date)
        scored = self.act_score_items(items, prefs)
        prefs = self.act_register_new_dishes(prefs, scored, items, event_date)
        description = self.act_build_description(scored)
        event_id = self.act_create_event(meal_slot, event_date, description)

        if simulated_response is not None:
            self._log(
                "A concrete RSVP was supplied for this run instead of polling Calendar — "
                "useful for exercising the decline-feedback path without waiting on a live invite.",
                "cli.simulated_response",
                {"value": simulated_response},
                f"Treating RSVP as '{simulated_response}'.",
            )
            response = simulated_response
        else:
            response = self.act_check_response(event_id)

        if response == "no":
            prefs = self.act_handle_decline(prefs, scored, event_id)
            event_id = None

        self.act_persist_preferences(prefs)
        return MealEvent(meal_slot=meal_slot, event_date=event_date, scored_items=scored, calendar_event_id=event_id, response=response)


def main() -> None:
    load_dotenv(BASE_DIR.parent / ".env")

    meal_slot = sys.argv[1] if len(sys.argv) > 1 else "lunch"
    event_date = date.fromisoformat(sys.argv[2]) if len(sys.argv) > 2 else date(2026, 9, 15)
    simulated_response = sys.argv[3] if len(sys.argv) > 3 else None

    agent = ReActMealAgent(
        user_id="u001",
        drive_skill=GoogleDriveSkill(),
        calendar_skill=GoogleCalendarSkill(),
        matcher=MenuPreferenceMatchingSkill(),
        fs=FilesystemTool(),
    )

    result = agent.run(meal_slot, event_date, simulated_response=simulated_response)

    print("\n=== Final MealEvent ===")
    print(result.model_dump_json(indent=2))


if __name__ == "__main__":
    main()
